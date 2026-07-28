"""
One-off script: Re-import assessment objectives on the production database.

Run this in Render Shell:
    python render_reimport_objectives.py

It uses the production DATABASE_URL environment variable automatically.
"""

import logging
import sys

import openpyxl

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")

# Import app modules (must run from the app directory, which it is on Render)
sys.path.insert(0, ".")
from app.database import SessionLocal
from app.importers.objectives import import_assessment_objectives
from app.models.assessment import AssessmentObjective

session = SessionLocal()

# 1. Delete all existing objectives
count = session.query(AssessmentObjective).delete()
session.commit()
print(f"Deleted {count} stale assessment objectives")

# 2. Re-import from the spreadsheet
wb = openpyxl.load_workbook(
    "secure-controls-framework-scf-2026-1.xlsx",
    data_only=True,
    read_only=False,
)
new_count = import_assessment_objectives(session, wb)
wb.close()
session.commit()
print(f"Imported {new_count} fresh assessment objectives")

# 3. Verify
from sqlalchemy import text
r = session.execute(text("SELECT COUNT(*) FROM assessment_objectives WHERE objective_code IS NULL"))
null_codes = r.scalar()
r = session.execute(text("SELECT COUNT(*) FROM assessment_objectives WHERE objective_text = 'SCF Created'"))
placeholders = r.scalar()
print(f"\nVerification:")
print(f"  Objectives with null codes:     {null_codes}  (should be 0)")
print(f"  'SCF Created' placeholders:     {placeholders}  (should be 0)")
print(f"  Total objectives:               {new_count}  (should be ~5,776)")

session.close()

if null_codes == 0 and placeholders == 0:
    print("\n✅ All good! The objectives are now properly imported.")
else:
    print("\n⚠️  Something didn't clean up fully — check the counts above.")
